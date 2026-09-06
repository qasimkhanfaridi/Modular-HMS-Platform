"""
Generate Word documents from LOOM replica project markdown pack.
"""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os
import re

DOCS_DIR = r"C:\Users\Faridi\Project\AAQSOLS\Modular-HMS-Platform\replica-project\docs"
OUTPUT_DIR = r"C:\Users\Faridi\Project\AAQSOLS\Modular-HMS-Platform\replica-project\word"

DOC_FILES = [
    ("00-Strategy-What-To-Do.md", "AAQSOLS-Replica-00-Strategy.docx"),
    ("01-BRD-Business-Requirements.md", "AAQSOLS-Replica-01-BRD.docx"),
    ("02-FRS-Functional-Requirements.md", "AAQSOLS-Replica-02-FRS.docx"),
    ("03-Video-Audio-Feature-Mapping.md", "AAQSOLS-Replica-03-Video-Audio-Mapping.docx"),
    ("04-Screen-Inventory-UI-Spec.md", "AAQSOLS-Replica-04-Screen-Inventory.docx"),
    ("05-User-Flows-Processes.md", "AAQSOLS-Replica-05-User-Flows.docx"),
    ("06-Proposal-Replica-Build.md", "AAQSOLS-Replica-06-Proposal.docx"),
    ("07-Demo-MVP-Roadmap.md", "AAQSOLS-Replica-07-Demo-Roadmap.docx"),
    ("08-Milestones-Development-Plan.md", "AAQSOLS-Replica-08-Milestones.docx"),
]


def set_doc_defaults(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def parse_table_row(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def is_table_separator(line):
    return bool(re.match(r"^\|[\s\-:|]+\|$", line.strip()))


def is_table_row(line):
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and not is_table_separator(s)


def add_table_from_rows(doc, headers, rows):
    if not headers:
        return
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        for p in table.rows[0].cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        for ci in range(len(headers)):
            val = row[ci] if ci < len(row) else ""
            table.rows[ri + 1].cells[ci].text = val
    doc.add_paragraph()


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def markdown_to_docx(md_path, doc):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    i = 0
    table_headers = None
    table_rows = []
    code_lines = []
    in_code = False

    def flush_table():
        nonlocal table_headers, table_rows
        if table_headers is not None:
            add_table_from_rows(doc, table_headers, table_rows)
            table_headers = None
            table_rows = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                flush_table()
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if is_table_row(stripped):
            cells = parse_table_row(stripped)
            if is_table_separator(stripped):
                i += 1
                continue
            if table_headers is None:
                table_headers = cells
            else:
                table_rows.append(cells)
            i += 1
            continue
        else:
            flush_table()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            doc.add_paragraph()
            i += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped.lstrip("#").strip()
            doc.add_heading(text, level=min(level, 4))
            i += 1
            continue

        if stripped.startswith("- [ ]") or stripped.startswith("- [x]") or stripped.startswith("- [X]"):
            text = stripped[5:].strip()
            doc.add_paragraph(text, style="List Bullet")
            i += 1
            continue

        if stripped.startswith("- ") or stripped.startswith("* "):
            doc.add_paragraph(stripped[2:].strip(), style="List Bullet")
            i += 1
            continue

        if re.match(r"^\d+\.\s", stripped):
            text = re.sub(r"^\d+\.\s", "", stripped)
            doc.add_paragraph(text, style="List Number")
            i += 1
            continue

        if stripped.startswith(">"):
            p = doc.add_paragraph(stripped.lstrip(">").strip())
            p.paragraph_format.left_indent = Inches(0.25)
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        doc.add_paragraph(stripped)
        i += 1

    flush_table()


def build_single(md_name, out_name):
    md_path = os.path.join(DOCS_DIR, md_name)
    doc = Document()
    set_doc_defaults(doc)
    markdown_to_docx(md_path, doc)
    out_path = os.path.join(OUTPUT_DIR, out_name)
    doc.save(out_path)
    return out_path


def build_combined():
    doc = Document()
    set_doc_defaults(doc)
    title = doc.add_heading("AAQSOLS Heart Clinic HMS — LOOM Replica Complete Pack", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "Business requirements, functional specs, video/audio mapping, UI inventory, "
        "flows, proposal, and demo roadmap. Prepared by AAQSOLS — September 2026."
    )
    doc.add_page_break()

    for md_name, _ in DOC_FILES:
        md_path = os.path.join(DOCS_DIR, md_name)
        markdown_to_docx(md_path, doc)
        doc.add_page_break()

    out_path = os.path.join(OUTPUT_DIR, "AAQSOLS-Replica-Complete-Pack.docx")
    doc.save(out_path)
    return out_path


def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    created = []
    for md_name, out_name in DOC_FILES:
        path = build_single(md_name, out_name)
        created.append(path)
        print(f"Created: {path}")
    combined = build_combined()
    print(f"Created: {combined}")
    print(f"\nDone — {len(created) + 1} Word files in {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
